import io
import csv
import uuid
from typing import List, Optional, Dict, Any
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from openpyxl import load_workbook
from config.settings import settings
from src.api.routes_auth import get_current_user, require_roles
from src.database.connection import get_db
from src.database.schemas import (
    ComplaintResponse, ComplaintAssignment, ComplaintStatusUpdate, ComplaintStatus, UserResponse, UserRole
)
from src.utils.geo_helpers import fleet_simulator, INITIAL_COMMUNITY_BINS
from src.utils.logger import get_logger

logger = get_logger("api.routes_admin")
router = APIRouter(prefix="/api/admin", tags=["Municipal Admin & Duty Allocation"])


@router.put("/complaints/{complaint_id}/assign", response_model=ComplaintResponse)
async def assign_complaint_duty(
    complaint_id: str,
    assignment: ComplaintAssignment,
    current_user: UserResponse = Depends(require_roles([UserRole.ADMIN])),
    db=Depends(get_db)
):
    """
    Municipal Officer assigns a citizen waste complaint to a sanitation staff member
    and assigns a municipal vehicle truck.
    """
    complaint = await db.complaints.find_one({"_id": complaint_id})
    if not complaint:
        raise HTTPException(status_code=404, detail="Complaint not found")

    now = settings.get_current_ist_time()
    update_data = {
        "status": ComplaintStatus.ASSIGNED.value,
        "assigned_to": assignment.model_dump(),
        "updated_at": now
    }

    await db.complaints.update_one({"_id": complaint_id}, {"$set": update_data})
    updated_doc = await db.complaints.find_one({"_id": complaint_id})
    
    logger.info(
        f"Complaint {complaint_id} assigned to staff {assignment.assigned_staff_name} by admin {current_user.id}"
    )
    return ComplaintResponse.model_validate(updated_doc)


@router.put("/complaints/{complaint_id}/status", response_model=ComplaintResponse)
async def update_complaint_status(
    complaint_id: str,
    status_update: ComplaintStatusUpdate,
    current_user: UserResponse = Depends(require_roles([UserRole.ADMIN, UserRole.STAFF])),
    db=Depends(get_db)
):
    """
    Updates complaint status (Pending -> Assigned -> In Progress -> Resolved)
    with resolution notes and proof photo.
    """
    complaint = await db.complaints.find_one({"_id": complaint_id})
    if not complaint:
        raise HTTPException(status_code=404, detail="Complaint not found")

    now = settings.get_current_ist_time()
    update_data = {
        "status": status_update.status.value,
        "updated_at": now
    }
    if status_update.resolution_notes:
        update_data["resolution_notes"] = status_update.resolution_notes
    if status_update.resolution_photo_url:
        update_data["resolution_photo_url"] = status_update.resolution_photo_url

    await db.complaints.update_one({"_id": complaint_id}, {"$set": update_data})
    updated_doc = await db.complaints.find_one({"_id": complaint_id})
    
    logger.info(f"Complaint {complaint_id} status updated to {status_update.status.value} by {current_user.id}")
    return ComplaintResponse.model_validate(updated_doc)


@router.get("/staff", response_model=List[UserResponse])
async def list_sanitation_staff(
    ward_no: Optional[str] = None,
    current_user: UserResponse = Depends(require_roles([UserRole.ADMIN])),
    db=Depends(get_db)
):
    """Lists sanitation staff members available for duty allocation."""
    query = {"role": UserRole.STAFF.value}
    if ward_no:
        query["ward_no"] = ward_no.upper()

    cursor = db.users.find(query)
    results = await cursor.to_list(length=100)
    return [UserResponse.model_validate(u) for u in results]


@router.get("/stats")
async def get_municipal_dashboard_stats(
    current_user: UserResponse = Depends(require_roles([UserRole.ADMIN])),
    db=Depends(get_db)
):
    """Returns real-time Swachh Bharat municipal metrics across wards."""
    total_complaints = await db.complaints.count_documents({})
    pending = await db.complaints.count_documents({"status": ComplaintStatus.PENDING.value})
    assigned = await db.complaints.count_documents({"status": ComplaintStatus.ASSIGNED.value})
    in_progress = await db.complaints.count_documents({"status": ComplaintStatus.IN_PROGRESS.value})
    resolved = await db.complaints.count_documents({"status": ComplaintStatus.RESOLVED.value})

    now_ist = settings.get_current_ist_time()
    date_str = now_ist.strftime("%Y-%m-%d")
    active_staff_count = await db.attendance.count_documents({"date_str": date_str, "status": "Active"})

    clearance_rate = round((resolved / total_complaints * 100.0), 1) if total_complaints > 0 else 100.0

    return {
        "ulb_name": settings.MUNICIPAL_NAME,
        "ulb_code": settings.MUNICIPAL_CODE,
        "timestamp": now_ist.isoformat(),
        "complaints": {
            "total": total_complaints,
            "pending": pending,
            "assigned": assigned,
            "in_progress": in_progress,
            "resolved": resolved,
            "clearance_rate_percent": clearance_rate
        },
        "fleet": {
            "active_trucks": len(fleet_simulator.state)
        },
        "workforce": {
            "active_on_duty_today": active_staff_count
        }
    }


REQUIRED_BIN_COLUMNS = ["ward_no", "landmark", "latitude", "longitude"]


@router.get("/bins", response_model=List[Dict[str, Any]])
async def list_community_bins(
    ward_no: Optional[str] = None,
    current_user: UserResponse = Depends(require_roles([UserRole.ADMIN])),
    db=Depends(get_db)
):
    """Lists all community waste bins currently stored in the database."""
    count = await db.bins.count_documents({})
    if count == 0:
        # First-run seed from the built-in defaults so the collection is never empty.
        await db.bins.insert_many([{**b, "_id": b["bin_id"]} for b in INITIAL_COMMUNITY_BINS])

    query = {}
    if ward_no:
        query["ward_no"] = ward_no.upper()

    cursor = db.bins.find(query)
    results = await cursor.to_list(length=500)
    for r in results:
        r.pop("_id", None)
    return results


@router.post("/bins/upload", status_code=status.HTTP_201_CREATED)
async def bulk_upload_bins_from_excel(
    file: UploadFile = File(...),
    current_user: UserResponse = Depends(require_roles([UserRole.ADMIN])),
    db=Depends(get_db)
):
    """
    Bulk-adds community waste bins from an uploaded Excel (.xlsx) sheet.
    Expected header row (case-insensitive, any column order):
      ward_no | landmark | latitude | longitude | wet_fill | dry_fill | hazardous_fill
    (wet_fill/dry_fill/hazardous_fill are optional 0-100 fill percentages, default 0)
    """
    filename_lower = file.filename.lower()
    if not filename_lower.endswith((".xlsx", ".xlsm", ".csv")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Please upload a valid Excel (.xlsx) or CSV file."
        )

    raw_content = await file.read()

    try:
        if filename_lower.endswith(".csv"):
            text = raw_content.decode("utf-8-sig")
            reader = csv.reader(io.StringIO(text))
            rows = [tuple(row) for row in reader]
        else:
            workbook = load_workbook(io.BytesIO(raw_content), data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not read the uploaded file: {str(e)}"
        )

    if not rows:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="The uploaded sheet is empty.")

    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    missing = [c for c in REQUIRED_BIN_COLUMNS if c not in header]
    if missing:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Missing required column(s): {', '.join(missing)}. "
                   f"Expected headers: ward_no, landmark, latitude, longitude "
                   f"(optional: wet_fill, dry_fill, hazardous_fill)."
        )
    col_idx = {name: header.index(name) for name in header if name}

    inserted, skipped, errors = 0, 0, []
    new_docs = []

    for row_num, row in enumerate(rows[1:], start=2):
        if row is None or all(cell is None for cell in row):
            continue
        try:
            ward_no = str(row[col_idx["ward_no"]]).strip().upper()
            landmark = str(row[col_idx["landmark"]]).strip()
            latitude = float(row[col_idx["latitude"]])
            longitude = float(row[col_idx["longitude"]])

            if not ward_no or not landmark:
                raise ValueError("ward_no and landmark cannot be empty")

            def fill(col_name):
                if col_name in col_idx and row[col_idx[col_name]] is not None:
                    val = int(float(row[col_idx[col_name]]))
                    return max(0, min(100, val))
                return 0

            bin_id = f"BIN-{uuid.uuid4().hex[:10].upper()}"
            new_docs.append({
                "_id": bin_id,
                "bin_id": bin_id,
                "ward_no": ward_no,
                "landmark": landmark,
                "latitude": latitude,
                "longitude": longitude,
                "bin_types": [
                    {"category": "Wet Waste", "color": "Green", "fill_level": fill("wet_fill")},
                    {"category": "Dry Waste", "color": "Blue", "fill_level": fill("dry_fill")},
                    {"category": "Hazardous/E-Waste", "color": "Black", "fill_level": fill("hazardous_fill")}
                ]
            })
            inserted += 1
        except Exception as e:
            skipped += 1
            errors.append(f"Row {row_num}: {str(e)}")

    if new_docs:
        await db.bins.insert_many(new_docs)

    logger.info(f"Admin {current_user.id} bulk-uploaded {inserted} bins via Excel ({skipped} skipped)")

    return {
        "inserted": inserted,
        "skipped": skipped,
        "errors": errors[:20]  # cap so a huge bad file doesn't return a massive payload
    }